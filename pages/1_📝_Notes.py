import time
import streamlit as st
from utils.database import (
    init_db, get_all_notes, get_notes_by_subject, get_all_subjects,
    add_note, delete_note, get_note, update_note, get_setting, toggle_favorite,
)
from utils.ai import generate_notes
from utils.helpers import (
    format_date, truncate_text, validate_note_input,
    render_sidebar, inject_custom_css, require_login,
    NOTE_TEMPLATES,
)

init_db()
require_login()

st.set_page_config(page_title="Notes - BrainDump", page_icon="📝", layout="wide")
inject_custom_css()
render_sidebar()

username = st.session_state.get("username", "")
show_confirmations = get_setting("show_confirmations")

for key, default in [
    ("notes_view", "list"), ("selected_note", None), ("editing_note", None),
    ("editor_title", ""), ("editor_subject", ""), ("editor_content", ""),
    ("editor_dirty", False), ("editor_last_saved", ""), ("editor_save_status", ""),
    ("last_autosave_time", 0),
]:
    if key not in st.session_state:
        st.session_state[key] = default


def _do_save(is_edit, note_id, title, subject, content):
    if is_edit:
        update_note(note_id, title, subject, content)
    else:
        new_id = add_note(username, title, subject, content)
        if new_id:
            st.session_state.editing_note = new_id
            st.session_state.notes_view = "edit"
    st.session_state.editor_dirty = False
    st.session_state.editor_last_saved = format_date(__import__("datetime").datetime.now())
    st.session_state.editor_save_status = "saved"


def _auto_save():
    from datetime import datetime
    title = st.session_state.get("editor_title", "")
    if not title.strip() or not st.session_state.editor_dirty:
        return
    now = time.time()
    if now - st.session_state.last_autosave_time < 30:
        return
    st.session_state.editor_save_status = "saving"
    _do_save(st.session_state.notes_view == "edit", st.session_state.editing_note,
             title, st.session_state.editor_subject, st.session_state.editor_content)
    st.session_state.last_autosave_time = now


st.title("📝 Notes")
st.caption("Create, edit, and organize your study notes")

# ═══════════════════════════════════════════════════════════
# LIST VIEW
# ═══════════════════════════════════════════════════════════
if st.session_state.notes_view == "list":
    subjects = get_all_subjects(username)
    filter_opts = ["All Subjects"] + subjects if subjects else ["All Subjects"]
    selected_subject = st.selectbox("Filter by subject", filter_opts)
    notes = get_all_notes(username) if selected_subject == "All Subjects" else get_notes_by_subject(username, selected_subject)

    show_fav = st.radio("Show", ["All Notes", "⭐ Favorites Only"], horizontal=True)
    if show_fav == "⭐ Favorites Only":
        notes = [n for n in notes if n.get("is_favorite")]

    search = st.text_input("🔍 Search notes", placeholder="Search by title or content...")
    if search:
        notes = [n for n in notes if search.lower() in n['title'].lower() or search.lower() in n['content'].lower()]

    if not notes:
        st.info("No notes found. Create your first one!")
    else:
        st.markdown(f"**{len(notes)} note(s)**")
        for note in notes:
            star = "⭐" if note.get("is_favorite") else "☆"
            st.markdown(f"""
            <div class="note-card">
                <div class="note-title">{star} {note['title']}</div>
                <span class="subject-tag">📚 {note['subject']}</span>
                <div class="note-content">{truncate_text(note['content'], 150)}</div>
                <div class="note-date">Created: {format_date(note['created_at'])}</div>
            </div>
            """, unsafe_allow_html=True)
            c1, c2, c3 = st.columns([1, 1, 4])
            with c1:
                if st.button("📖 Open", key=f"open_{note['id']}", help="View note"):
                    st.session_state.notes_view = "view"
                    st.session_state.selected_note = note['id']
                    st.rerun()
            with c2:
                btn_label = "⭐" if not note.get("is_favorite") else "★"
                if st.button(btn_label, key=f"fav_{note['id']}", help="Toggle favorite"):
                    toggle_favorite(note['id']); st.rerun()
            # Delete with confirmation
            if st.button("🗑️ Delete", key=f"del_ask_{note['id']}", help="Delete"):
                if show_confirmations:
                    st.session_state[f"cd_{note['id']}"] = True; st.rerun()
                else:
                    delete_note(note['id']); st.toast("Deleted", icon="🗑️"); st.rerun()
            if show_confirmations and st.session_state.get(f"cd_{note['id']}"):
                st.warning(f"Delete **{note['title']}**?")
                dcc1, dcc2 = st.columns(2)
                with dcc1:
                    if st.button("✅ Yes", key=f"dy_{note['id']}", type="primary"):
                        delete_note(note['id'])
                        st.session_state.pop(f"cd_{note['id']}", None)
                        st.toast("Deleted!", icon="🗑️"); st.rerun()
                with dcc2:
                    if st.button("❌ No", key=f"dn_{note['id']}"):
                        st.session_state.pop(f"cd_{note['id']}", None); st.rerun()

    st.markdown("---")
    if st.button("📝 Create New Note", type="primary"):
        st.session_state.notes_view = "create"
        st.session_state.editing_note = None
        st.session_state.editor_title = ""; st.session_state.editor_subject = ""; st.session_state.editor_content = ""
        st.session_state.editor_dirty = False; st.session_state.editor_save_status = ""
        st.rerun()

    # ── File Upload ──────────────────────────────────────────
    st.markdown("---")
    st.markdown("### 📄 Upload File")
    uploaded = st.file_uploader("Upload document", type=["pdf", "txt", "docx"], label_visibility="collapsed")
    if uploaded:
        ext = uploaded.name.rsplit(".", 1)[-1].lower()
        try:
            if ext == "txt":
                content = uploaded.read().decode("utf-8", errors="replace")
            elif ext == "pdf":
                import pdfplumber
                with pdfplumber.open(uploaded) as pdf:
                    content = "\n".join([p.extract_text() or "" for p in pdf.pages])
                from utils.ai import summarize_notes
                cleaned, err = summarize_notes(content)
                if not err and cleaned:
                    content = cleaned
            elif ext == "docx":
                import docx
                doc = docx.Document(uploaded)
                content = "\n".join([p.text for p in doc.paragraphs])
            else:
                content = ""
            if content.strip():
                title = uploaded.name.rsplit(".", 1)[0]
                nid = add_note(username, title, "Uploaded", content)
                if nid:
                    st.toast(f"Created from {uploaded.name}!", icon="✅")
                    st.session_state.notes_view = "view"; st.session_state.selected_note = nid; st.rerun()
            else:
                st.error("No text extracted from file")
        except Exception as e:
            st.error(f"Error: {e}")

    # ── AI Notes Generator ───────────────────────────────────
    st.markdown("---")
    st.markdown("### 🤖 Generate AI Notes")
    ai_c1, ai_c2 = st.columns(2)
    with ai_c1:
        ai_subj = st.selectbox("Subject", ["Maths", "Geography", "History", "Economics", "Software Engineering"], key="ai_subj")
        ai_topic = st.text_input("Topic (optional)", placeholder="Leave blank for random", key="ai_topic")
    with ai_c2:
        ai_diff = st.selectbox("Difficulty", ["Beginner", "Intermediate", "Advanced"], key="ai_diff")
        ai_len = st.selectbox("Length", ["Short", "Medium", "Long"], index=1, key="ai_len")
    if st.button("✨ Generate Notes", type="primary", key="ai_gen"):
        content, err = generate_notes(ai_subj, ai_topic, ai_diff, ai_len)
        if err:
            st.error(err)
        else:
            title = "AI Generated Notes"
            for line in content.split("\n"):
                if line.startswith("# ") and not line.startswith("## "):
                    title = line[2:].strip(); break
            nid = add_note(username, title, ai_subj, content)
            if nid:
                st.toast("Notes generated!", icon="✅")
                st.session_state.notes_view = "view"; st.session_state.selected_note = nid; st.rerun()
            else:
                st.error("Failed to save")

# ═══════════════════════════════════════════════════════════
# CREATE / EDIT VIEW
# ═══════════════════════════════════════════════════════════
elif st.session_state.notes_view in ("create", "edit"):
    from datetime import datetime
    if st.button("← Back to Notes"):
        if st.session_state.editor_dirty and st.session_state.editor_title.strip():
            _do_save(st.session_state.notes_view == "edit", st.session_state.editing_note,
                     st.session_state.editor_title, st.session_state.editor_subject, st.session_state.editor_content)
        st.session_state.notes_view = "list"; st.session_state.editing_note = None; st.rerun()

    st.markdown("---")
    is_edit = st.session_state.notes_view == "edit"
    st.markdown(f"### {'✏️ Edit Note' if is_edit else '📝 Create New Note'}")

    if is_edit and st.session_state.editing_note and not st.session_state.editor_title:
        existing = get_note(st.session_state.editing_note)
        if existing:
            st.session_state.editor_title = existing['title']
            st.session_state.editor_subject = existing['subject']
            st.session_state.editor_content = existing['content']

    if not is_edit:
        tmpl = st.selectbox("📋 Use Template", list(NOTE_TEMPLATES.keys()), key="tmpl_sel")
        if st.button("📋 Apply Template"):
            st.session_state.editor_content = NOTE_TEMPLATES[tmpl]; st.session_state.editor_dirty = True; st.rerun()

    new_title = st.text_input("Title", value=st.session_state.editor_title, placeholder="Enter note title...")
    if new_title != st.session_state.editor_title:
        st.session_state.editor_title = new_title; st.session_state.editor_dirty = True

    existing_subjects = get_all_subjects(username)
    opts = existing_subjects + ["+ Add new subject"]
    current_subj = st.session_state.editor_subject
    subj_idx = existing_subjects.index(current_subj) if current_subj in existing_subjects else 0
    choice = st.selectbox("Subject", opts, index=subj_idx if existing_subjects else 0)
    new_subject = st.text_input("New subject", placeholder="Subject...") if choice == "+ Add new subject" else choice
    if new_subject != st.session_state.editor_subject:
        st.session_state.editor_subject = new_subject; st.session_state.editor_dirty = True

    st.markdown('<div class="toolbar-hint">Supports Markdown: **bold**, *italic*, # heading, - bullet points</div>', unsafe_allow_html=True)
    new_content = st.text_area("Content", value=st.session_state.editor_content, height=500, placeholder="Write your notes here...")
    if new_content != st.session_state.editor_content:
        st.session_state.editor_content = new_content; st.session_state.editor_dirty = True

    words = len(new_content.split()) if new_content.strip() else 0
    st.markdown(f'<div class="save-status">{words} words · {len(new_content)} chars</div>', unsafe_allow_html=True)
    _auto_save()

    status = st.session_state.editor_save_status
    if status == "saved" and st.session_state.editor_last_saved:
        st.markdown(f'<div class="save-status">Saved ✓ ({st.session_state.editor_last_saved})</div>', unsafe_allow_html=True)
    elif st.session_state.editor_dirty:
        st.markdown('<div class="save-status">Unsaved changes</div>', unsafe_allow_html=True)

    col1, col2 = st.columns([1, 5])
    with col1:
        if st.button("💾 Save", type="primary"):
            errs = validate_note_input(st.session_state.editor_title, st.session_state.editor_subject, st.session_state.editor_content)
            if errs:
                for e in errs: st.error(e)
            else:
                _do_save(is_edit, st.session_state.editing_note, st.session_state.editor_title, st.session_state.editor_subject, st.session_state.editor_content)
                st.toast("Saved!", icon="✅")
                st.session_state.notes_view = "list"; st.session_state.editing_note = None
                st.session_state.editor_title = ""; st.session_state.editor_subject = ""; st.session_state.editor_content = ""
                st.rerun()
    with col2:
        if st.button("❌ Cancel"):
            st.session_state.notes_view = "list"; st.session_state.editing_note = None
            st.session_state.editor_title = ""; st.session_state.editor_subject = ""; st.session_state.editor_content = ""
            st.session_state.editor_dirty = False; st.rerun()

# ═══════════════════════════════════════════════════════════
# VIEW SINGLE NOTE
# ═══════════════════════════════════════════════════════════
elif st.session_state.notes_view == "view":
    if st.button("← Back to Notes"):
        st.session_state.notes_view = "list"; st.session_state.selected_note = None; st.rerun()

    note = get_note(st.session_state.selected_note)
    if note:
        st.markdown("---")
        vc1, vc2 = st.columns([5, 1])
        with vc1:
            star = "⭐" if note.get("is_favorite") else "☆"
            st.markdown(f"## {star} {note['title']}")
        with vc2:
            if st.button("✏️ Edit"):
                st.session_state.notes_view = "edit"; st.session_state.editing_note = note['id']
                st.session_state.editor_title = ""; st.session_state.editor_subject = ""; st.session_state.editor_content = ""
                st.rerun()

        st.markdown(f"**📚 {note['subject']}** · Created: {format_date(note['created_at'])}")
        if note.get('updated_at') and note['updated_at'] != note['created_at']:
            st.caption(f"Last edited: {format_date(note['updated_at'])}")
        st.markdown("---")
        st.markdown(note['content'])

        # Export
        st.markdown("---")
        st.markdown("### 📥 Export")
        ec1, ec2 = st.columns(2)
        with ec1:
            txt = f"{note['title']}\nSubject: {note['subject']}\n\n{note['content']}"
            st.download_button("📄 Download .txt", txt, file_name=f"{note['title']}.txt", mime="text/plain", use_container_width=True)
        with ec2:
            try:
                from fpdf import FPDF
                pdf = FPDF()
                pdf.add_page()
                pdf.set_font("Helvetica", "B", 16)
                pdf.cell(0, 10, note['title'].encode('latin-1', 'replace').decode('latin-1'), new_x="LMARGIN", new_y="NEXT")
                pdf.set_font("Helvetica", "", 11)
                pdf.cell(0, 8, f"Subject: {note['subject']}", new_x="LMARGIN", new_y="NEXT")
                pdf.ln(5)
                pdf.multi_cell(0, 6, note['content'].encode('latin-1', 'replace').decode('latin-1'))
                st.download_button("📑 Download .pdf", bytes(pdf.output()), file_name=f"{note['title']}.pdf", mime="application/pdf", use_container_width=True)
            except Exception:
                st.caption("Install fpdf2 for PDF export")
    else:
        st.error("Note not found")
        st.session_state.notes_view = "list"
